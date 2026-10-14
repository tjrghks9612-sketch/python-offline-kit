"""Verify the delivered interpreter and real workflows without network access."""

from __future__ import annotations

import argparse
import importlib
import importlib.metadata as md
import io
import ipaddress
import json
import os
import platform
import re
import socket
import subprocess
import sys
import traceback
from datetime import datetime, timezone
from pathlib import Path


def no_network(*args: object, **kwargs: object) -> None:
    raise OSError("Verification blocks network access")


def is_loopback(host: object) -> bool:
    try:
        return ipaddress.ip_address(str(host)).is_loopback
    except ValueError:
        return False


original_connect = socket.socket.connect
original_connect_ex = socket.socket.connect_ex
original_create_connection = socket.create_connection
original_getaddrinfo = socket.getaddrinfo


def local_connect(sock, address):
    if not isinstance(address, tuple) or not is_loopback(address[0]):
        no_network()
    return original_connect(sock, address)


def local_connect_ex(sock, address):
    if not isinstance(address, tuple) or not is_loopback(address[0]):
        no_network()
    return original_connect_ex(sock, address)


def local_create_connection(address, *args, **kwargs):
    if not is_loopback(address[0]):
        no_network()
    return original_create_connection(address, *args, **kwargs)


def local_getaddrinfo(host, *args, **kwargs):
    if not is_loopback(host):
        no_network()
    return original_getaddrinfo(host, *args, **kwargs)


socket.socket.connect = local_connect
socket.socket.connect_ex = local_connect_ex
socket.create_connection = local_create_connection
socket.getaddrinfo = local_getaddrinfo


def audit(event: str, args: tuple) -> None:
    if event == "socket.connect" and not is_loopback(args[1][0]):
        no_network()
    if event == "socket.getaddrinfo" and not is_loopback(args[0]):
        no_network()
    if event == "socket.sendto" and not is_loopback(args[-1][0]):
        no_network()


sys.addaudithook(audit)

IMPORTS = {
    "numpy": "numpy",
    "pandas": "pandas",
    "polars": "polars",
    "pyarrow": "pyarrow",
    "scipy": "scipy",
    "scikit-learn": "sklearn",
    "statsmodels": "statsmodels.api",
    "matplotlib": "matplotlib",
    "seaborn": "seaborn",
    "plotly": "plotly",
    "duckdb": "duckdb",
    "sqlalchemy": "sqlalchemy",
    "pyodbc": "pyodbc",
    "openpyxl": "openpyxl",
    "xlsxwriter": "xlsxwriter",
    "xlrd": "xlrd",
    "python-docx": "docx",
    "python-pptx": "pptx",
    "pypdf": "pypdf",
    "pdfplumber": "pdfplumber",
    "reportlab": "reportlab",
    "pillow": "PIL.Image",
    "requests": "requests",
    "httpx": "httpx",
    "beautifulsoup4": "bs4",
    "lxml": "lxml.etree",
    "fastapi": "fastapi",
    "uvicorn": "uvicorn",
    "python-multipart": "python_multipart",
    "streamlit": "streamlit",
    "jinja2": "jinja2",
    "pyyaml": "yaml",
    "python-dotenv": "dotenv",
    "pydantic": "pydantic",
    "cryptography": "cryptography",
    "jupyterlab": "jupyterlab",
    "ipykernel": "ipykernel",
    "ipywidgets": "ipywidgets",
    "nbconvert": "nbconvert",
    "rich": "rich",
    "tqdm": "tqdm",
    "typer": "typer",
    "tenacity": "tenacity",
    "rapidfuzz": "rapidfuzz",
    "psutil": "psutil",
    "pywin32": "win32api",
    "pytest": "pytest",
    "pytest-cov": "pytest_cov",
    "ruff": "ruff",
    "black": "black",
    "mypy": "mypy",
    "build": "build",
}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, required=True)
    args = parser.parse_args()
    root = args.root.resolve()
    output = root / "verification-output"
    output.mkdir(exist_ok=True)
    os.environ["MPLBACKEND"] = "Agg"
    os.environ["MPLCONFIGDIR"] = str(output / "matplotlib-cache")
    os.environ["STREAMLIT_BROWSER_GATHER_USAGE_STATS"] = "false"
    os.environ["PYTHONUTF8"] = "1"
    manifest = json.loads((root / "bundle-manifest.json").read_text(encoding="utf-8"))
    results = []

    def check(name, function):
        try:
            detail = function()
            results.append({"name": name, "status": "PASS", "detail": detail})
            print(f"PASS {name}", flush=True)
        except Exception as exc:
            results.append(
                {
                    "name": name,
                    "status": "FAIL",
                    "error": str(exc),
                    "traceback": traceback.format_exc(),
                }
            )
            print(f"FAIL {name}: {exc}", flush=True)

    def command(arguments):
        completed = subprocess.run(
            arguments,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
            cwd=output,
        )
        if completed.returncode:
            raise RuntimeError((completed.stdout + completed.stderr)[-4000:])
        return completed.stdout.strip()[-1000:]

    def runtime():
        import struct
        import sysconfig

        assert sys.platform == "win32" and struct.calcsize("P") == 8
        assert sysconfig.get_platform() == "win-amd64"
        assert platform.python_version() == manifest["python"]["version"]
        assert Path(sys.prefix).resolve() == (root / ".venv").resolve(), sys.prefix
        assert sys.prefix != sys.base_prefix
        assert not sysconfig.get_config_var("Py_GIL_DISABLED")
        return {
            "python": platform.python_version(),
            "platform": platform.platform(),
            "architecture": "x86-64",
        }

    def locked_packages():
        expected = {}
        for line in (root / "requirements.lock").read_text(encoding="utf-8").splitlines():
            if not line or line.startswith("#"):
                continue
            match = re.fullmatch(r"([A-Za-z0-9_.-]+)==([^ ]+) --hash=sha256:[0-9a-f]{64}", line)
            assert match, f"Invalid lock entry: {line}"
            name, version = match.groups()
            actual = md.version(name)
            assert actual == version, f"{name}: {actual} != {version}"
            expected[name] = version
        assert len(expected) == manifest["wheel_count"]
        return {"packages_checked": len(expected)}

    def imports():
        direct = [
            x.strip().lower().replace("_", "-")
            for x in (root / "requirements.in").read_text(encoding="utf-8").splitlines()
            if x.strip() and not x.lstrip().startswith("#")
        ]
        assert set(direct) == set(IMPORTS), "Update the import map when direct packages change"
        for name in direct:
            importlib.import_module(IMPORTS[name])
        return {"direct_imports_checked": len(direct)}

    def numeric():
        import numpy as np
        from scipy.linalg import solve

        solution = solve(np.array([[3.0, 1.0], [1.0, 2.0]]), np.array([9.0, 8.0]))
        assert np.allclose(solution, [2, 3])

    def excel():
        import pandas as pd

        frame = pd.DataFrame({"부서": ["영업", "기획"], "매출": [120, 80]})
        path = output / "sample.xlsx"
        frame.to_excel(path, index=False, engine="openpyxl")
        pd.testing.assert_frame_equal(pd.read_excel(path, engine="openpyxl"), frame)

    def xlsxwriter():
        import openpyxl
        import xlsxwriter as xw

        path = output / "xlsxwriter-sample.xlsx"
        with xw.Workbook(path) as workbook:
            sheet = workbook.add_worksheet("결과")
            sheet.write(0, 0, "오프라인 설치 성공")
            sheet.write_number(1, 0, 42)
        workbook = openpyxl.load_workbook(path)
        assert workbook["결과"]["A2"].value == 42
        workbook.close()

    def columnar():
        import duckdb
        import pandas as pd
        import polars as pl

        path = output / "sample.parquet"
        pd.DataFrame({"value": [1, 2, 3]}).to_parquet(path, engine="pyarrow")
        assert pl.read_parquet(path)["value"].sum() == 6
        with duckdb.connect() as connection:
            assert (
                connection.execute(
                    "SELECT sum(value) FROM read_parquet(?)", [str(path)]
                ).fetchone()[0]
                == 6
            )

    def machine_learning():
        import numpy as np
        from sklearn.linear_model import LinearRegression

        model = LinearRegression().fit(np.array([[1], [2], [3]]), [3, 5, 7])
        assert np.isclose(model.predict([[4]])[0], 9)

    def statistics():
        import numpy as np
        import statsmodels.api as sm

        result = sm.OLS([3, 5, 7, 9], sm.add_constant([1, 2, 3, 4])).fit()
        assert np.allclose(result.params, [1, 2])

    def chart():
        import matplotlib.pyplot as plt
        from matplotlib.font_manager import FontProperties
        from PIL import Image

        font_path = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "malgun.ttf"
        font = FontProperties(fname=font_path) if font_path.exists() else None
        fig, ax = plt.subplots(figsize=(5, 3))
        ax.bar(["A", "B", "C"], [120, 80, 100])
        ax.set_title("오프라인 매출 차트" if font else "Offline sales chart", fontproperties=font)
        fig.tight_layout()
        path = output / "chart.png"
        fig.savefig(path, dpi=120)
        plt.close(fig)
        with Image.open(path) as image:
            assert image.width == 600 and image.height == 360
        return {"korean_font_available": bool(font)}

    def interactive_chart():
        import plotly.graph_objects as go

        path = output / "chart.html"
        go.Figure(go.Bar(x=["A", "B"], y=[1, 2])).write_html(
            path, include_plotlyjs=True, full_html=True
        )
        assert path.stat().st_size > 1_000_000

    def word():
        from docx import Document

        path = output / "sample.docx"
        doc = Document()
        doc.add_heading("오프라인 검증", 0)
        doc.add_paragraph("파이썬 문서 생성 확인")
        doc.save(path)
        assert "파이썬" in Document(path).paragraphs[1].text

    def powerpoint():
        from pptx import Presentation

        path = output / "sample.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "오프라인 검증"
        presentation.save(path)
        assert Presentation(path).slides[0].shapes.title.text == "오프라인 검증"

    def pdf():
        import pdfplumber
        import pypdfium2
        from pypdf import PdfReader
        from reportlab.pdfgen.canvas import Canvas

        path = output / "sample.pdf"
        canvas = Canvas(str(path))
        canvas.drawString(72, 720, "Offline PDF verification")
        canvas.save()
        assert "Offline PDF" in PdfReader(path).pages[0].extract_text()
        with pdfplumber.open(path) as pdf_document:
            assert "Offline PDF" in pdf_document.pages[0].extract_text()
        pdf_document = pypdfium2.PdfDocument(path)
        try:
            bitmap = pdf_document[0].render(scale=1)
            bitmap.to_pil().save(output / "pdf-preview.png")
        finally:
            pdf_document.close()

    def database():
        from sqlalchemy import create_engine, text

        engine = create_engine("sqlite+pysqlite:///:memory:")
        with engine.begin() as connection:
            assert connection.execute(text("SELECT 40 + 2")).scalar_one() == 42
        engine.dispose()

    def odbc():
        import pyodbc

        return {
            "installed_odbc_drivers": pyodbc.drivers(),
            "note": "Database vendor drivers and credentials are supplied separately.",
        }

    def web_application():
        from fastapi import FastAPI
        from fastapi.testclient import TestClient

        app = FastAPI()

        @app.get("/health")
        def health():
            return {"status": "ok", "message": "검증 성공"}

        with TestClient(app) as client:
            response = client.get("/health")
            assert response.status_code == 200 and response.json()["message"] == "검증 성공"

    def html():
        from bs4 import BeautifulSoup

        assert BeautifulSoup("<h1>검증</h1>", "lxml").h1.text == "검증"

    def configuration():
        import yaml
        from dotenv import dotenv_values
        from jinja2 import Template
        from pydantic import BaseModel

        class Settings(BaseModel):
            count: int

        assert Settings(count="42").count == 42
        assert yaml.safe_load("count: 42")["count"] == 42
        assert dotenv_values(stream=io.StringIO("AX_TEST=local"))["AX_TEST"] == "local"
        assert Template("Hello {{ name }}").render(name="AX") == "Hello AX"

    def encryption():
        from cryptography.fernet import Fernet

        cipher = Fernet(Fernet.generate_key())
        assert cipher.decrypt(cipher.encrypt(b"offline")) == b"offline"

    def notebook():
        import ipywidgets
        import nbformat
        from nbconvert import HTMLExporter

        notebook = nbformat.v4.new_notebook(
            cells=[nbformat.v4.new_markdown_cell("# Offline verification")]
        )
        nbformat.write(notebook, output / "sample.ipynb")
        html_body, _ = HTMLExporter().from_notebook_node(notebook)
        assert "Offline verification" in html_body
        (output / "notebook.html").write_text(html_body, encoding="utf-8")
        assert ipywidgets.IntSlider(value=7).value == 7
        return command([sys.executable, "-I", "-m", "jupyterlab", "--version"])

    def windows_automation():
        import pythoncom
        import win32api

        assert win32api.GetCurrentProcessId() == os.getpid()
        pythoncom.CoInitialize()
        pythoncom.CoUninitialize()

    def developer_tools():
        sample = output / "test_sample.py"
        sample.write_text(
            "def test_arithmetic() -> None:\n    assert 20 + 22 == 42\n", encoding="utf-8"
        )
        command([sys.executable, "-I", "-m", "pytest", "-q", "-p", "no:cacheprovider", str(sample)])
        command([str(root / ".venv/Scripts/ruff.exe"), "check", str(sample)])
        command([sys.executable, "-I", "-m", "black", "--check", str(sample)])
        return command(
            [
                sys.executable,
                "-I",
                "-m",
                "mypy",
                "--no-incremental",
                "--cache-dir",
                str(output / "mypy-cache"),
                str(sample),
            ]
        )

    def network_guard():
        try:
            socket.create_connection(("example.com", 443))
        except OSError as exc:
            assert "Verification blocks network" in str(exc)
        else:
            raise AssertionError("Network guard is inactive")

    for name, function in [
        ("interpreter", runtime),
        ("all_pinned_versions", locked_packages),
        ("all_direct_imports", imports),
        ("numpy_scipy_calculation", numeric),
        ("korean_excel_roundtrip", excel),
        ("xlsxwriter_roundtrip", xlsxwriter),
        ("parquet_polars_duckdb", columnar),
        ("scikit_learn_fit", machine_learning),
        ("statsmodels_fit", statistics),
        ("png_chart", chart),
        ("standalone_plotly_html", interactive_chart),
        ("word_roundtrip", word),
        ("powerpoint_roundtrip", powerpoint),
        ("pdf_create_extract_render", pdf),
        ("sqlalchemy_sqlite", database),
        ("odbc_driver_enumeration", odbc),
        ("fastapi_in_process_request", web_application),
        ("html_parsing", html),
        ("configuration_validation", configuration),
        ("cryptography_roundtrip", encryption),
        ("notebook_html_and_jupyterlab", notebook),
        ("windows_com", windows_automation),
        ("pytest_ruff_black_mypy", developer_tools),
        ("network_guard", network_guard),
    ]:
        check(name, function)
    failed = sum(r["status"] == "FAIL" for r in results)
    report = {
        "status": "PASS" if failed == 0 else "FAIL",
        "checked_at": datetime.now(timezone.utc).isoformat(),
        "network": "External connections and DNS are blocked; loopback is allowed for Windows asyncio internals. No external service calls are made.",
        "tests": len(results),
        "passed": len(results) - failed,
        "failed": failed,
        "results": results,
    }
    (output / "verification.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(
        f"Verification: {report['passed']}/{report['tests']} passed. Report: {output / 'verification.json'}"
    )
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
